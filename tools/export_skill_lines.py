"""Export each legendary lord's unique skill line (lane 1) to an .xlsx.

Reads the TSVs tools/check_skill_trees.py leaves in .skilltree_cache/extract, so run that
first (or after any pack edit) to refresh them.

    py tools/check_skill_trees.py && py tools/export_skill_lines.py
"""
import collections, glob, os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # tools/ -> workspace root
EX = os.path.join(ROOT, ".skilltree_cache", "extract")
OUT = os.path.join(ROOT, "docs", "HASHUT_LL_SKILL_LINES.xlsx")

def tsv(folder, root="db"):
    rows = []
    for p in sorted(glob.glob(os.path.join(EX, root, folder, "*.tsv"))):
        lines = open(p, encoding="utf-8").read().splitlines()
        cols = lines[0].split("\t")
        for ln in lines[1:]:
            if ln.startswith("#") or not ln.strip():
                continue
            v = ln.split("\t")
            v += [""] * (len(cols) - len(v))
            rows.append(dict(zip(cols, v)))
    return rows

loc = {}
for p in glob.glob(os.path.join(EX, "text", "*.tsv")):
    for ln in open(p, encoding="utf-8").read().splitlines()[1:]:
        f = ln.split("\t")
        if not ln.startswith("#") and len(f) >= 2:
            loc[f[0]] = f[1]

nodes = {n["key"]: n for n in tsv("character_skill_nodes_tables")}
skills = {s["key"]: s for s in tsv("character_skills_tables")}
sets = tsv("character_skill_node_sets_tables")
items = collections.defaultdict(set)
for it in tsv("character_skill_node_set_items_tables"):
    items[it["set"]].add(it["item"])
eff = collections.defaultdict(list)
for e in tsv("character_skill_level_to_effects_junctions_tables"):
    eff[e["character_skill_key"]].append(e)
ability = {r["effect"]: r["unit_ability"]
           for r in tsv("effect_bonus_value_unit_ability_junctions_tables")}

def n(x):
    try:
        return int(float(x))
    except ValueError:
        return 0

HEAD = ["Lord", "Slot", "Skill", "Unlocks at rank", "Description",
        "Effects", "Grants ability", "Node key", "Skill key", "Icon"]

wb = Workbook()
ws = wb.active
ws.title = "Unique skill lines"
ws.append(HEAD)

band = PatternFill("solid", fgColor="F2F2F2")
for s in sets:
    lord = loc.get("land_units_onscreen_name_" + s["agent_subtype_key"], s["agent_subtype_key"])
    line = sorted((k for k in items[s["key"]] if k in nodes and nodes[k]["indent"] == "1"),
                  key=lambda k: n(nodes[k]["tier"]))
    for k in line:
        sk = nodes[k]["character_skill_key"]
        row = skills.get(sk, {})
        effects = "\n".join(
            "%s  %s  (%s)" % (e["value"].rstrip("0").rstrip(".") or "0", e["effect_key"], e["effect_scope"])
            for e in eff[sk])
        grants = "\n".join(ability[e["effect_key"]] for e in eff[sk] if e["effect_key"] in ability)
        ws.append([lord,
                   n(nodes[k]["tier"]),
                   loc.get("character_skills_localised_name_" + sk, row.get("localised_name", "")),
                   n(row.get("unlocked_at_rank")),
                   loc.get("character_skills_localised_description_" + sk,
                           row.get("localised_description", "")),
                   effects, grants, k, sk, row.get("image_path", "")])

# formatting
for c in range(1, len(HEAD) + 1):
    ws.cell(1, c).font = Font(bold=True)
    ws.cell(1, c).fill = PatternFill("solid", fgColor="D9D9D9")
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
for col, w in zip("ABCDEFGHIJ", (20, 6, 30, 10, 60, 70, 30, 24, 24, 34)):
    ws.column_dimensions[col].width = w
for r in range(2, ws.max_row + 1):
    for c in range(1, len(HEAD) + 1):
        ws.cell(r, c).alignment = Alignment(vertical="top",
                                            wrap_text=c in (3, 5, 6, 7))
    if (r - 2) // 6 % 2:                    # band every other lord's six rows
        for c in range(1, len(HEAD) + 1):
            ws.cell(r, c).fill = band

wb.save(OUT)
print("%s  -  %d skills across %d lords" % (OUT, ws.max_row - 1, len(sets)))
